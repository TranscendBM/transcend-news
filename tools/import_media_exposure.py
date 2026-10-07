"""Import manually verified media exposure workbooks into Firestore.

Safety defaults:
  * Running without ``--copy`` is read-only and only prints a dry-run report.
  * Source workbooks are never modified.
  * Public records never contain internal notes, source filenames, or row numbers.
  * Firestore writes require an explicit project id and approval flag.
  * The importer never deletes documents.

The public ``media_exposure`` collection is readable by the dashboard.  Raw
provenance and internal notes are written to ``media_exposure_private``, which
is intentionally covered by Firestore Rules' default deny rule.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from zoneinfo import ZoneInfo


TAIPEI = ZoneInfo("Asia/Taipei")
DEST_PROJECT = "transcend-news-tbm"
PUBLIC_COLLECTION = "media_exposure"
PRIVATE_COLLECTION = "media_exposure_private"
MAX_RECORDS_PER_BATCH = 200  # two writes per record; stays below 500 ops

HEADER_ALIASES = {
    "日期": "date",
    "媒體": "media",
    "記者": "reporter",
    "新聞標題": "title",
    "標題": "title",
    "連結": "link",
    "網址": "link",
    "備註": "note",
    "內部備註": "note",
}

MEDIA_ALIASES = {
    "工商": "工商時報",
    "工商時報": "工商時報",
    "時報新聞": "時報資訊",
    "時報記者": "時報資訊",
    "時報資訊": "時報資訊",
    "經濟日報記者": "經濟日報",
    "聯合新聞網": "聯合報",
}

TRACKING_QUERY_KEYS = {
    "fbclid", "gclid", "mc_cid", "mc_eid", "ref", "source",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_media(value: Any) -> str:
    name = _text(value)
    return MEDIA_ALIASES.get(name, name)


def normalize_title(value: Any) -> str:
    return _text(value)


def parse_exposure_date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        raw = value
    elif isinstance(value, date):
        raw = datetime(value.year, value.month, value.day)
    else:
        text = _text(value)
        if not text:
            return None
        raw = None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%m/%d/%Y"):
            try:
                raw = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if raw is None:
            return None
    if raw.tzinfo is None:
        return raw.replace(tzinfo=TAIPEI)
    return raw.astimezone(TAIPEI)


def canonicalize_url(value: Any) -> str:
    raw = _text(value)
    if not raw:
        return ""
    if any(ch.isspace() or ord(ch) < 32 for ch in raw):
        return ""
    if "://" not in raw:
        raw = "https://" + raw
    try:
        parts = urlsplit(raw)
    except ValueError:
        return ""
    if parts.scheme.lower() not in {"http", "https"} or not parts.hostname:
        return ""
    if parts.username or parts.password:
        return ""
    host = parts.hostname.lower()
    port = parts.port
    if port and not ((parts.scheme.lower() == "http" and port == 80)
                     or (parts.scheme.lower() == "https" and port == 443)):
        host = f"{host}:{port}"
    path = parts.path or "/"
    query = [
        (key, val) for key, val in parse_qsl(parts.query, keep_blank_values=True)
        if not key.lower().startswith("utm_") and key.lower() not in TRACKING_QUERY_KEYS
    ]
    return urlunsplit((parts.scheme.lower(), host, path, urlencode(query), ""))


def infer_exposure_type(url: str) -> str:
    host = (urlsplit(url).hostname or "").lower()
    if host in {"youtube.com", "www.youtube.com", "youtu.be"}:
        return "video"
    if host.endswith("facebook.com") or host.endswith("instagram.com"):
        return "social"
    return "online"


def exposure_id(canonical_url: str, exposure_date: datetime, media: str, title: str) -> str:
    identity = canonical_url or "|".join([
        exposure_date.date().isoformat(), media.casefold(), title.casefold(),
    ])
    return "exp_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:32]


@dataclass
class ExposureRecord:
    doc_id: str
    public: dict[str, Any]
    private: dict[str, Any]
    sources: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class ParseReport:
    rows_seen: int = 0
    valid_rows: int = 0
    invalid_rows: list[dict[str, Any]] = field(default_factory=list)
    duplicate_rows: int = 0
    records: dict[str, ExposureRecord] = field(default_factory=dict)

    def summary(self) -> dict[str, Any]:
        media = {record.public["mediaName"] for record in self.records.values()}
        types: dict[str, int] = {}
        for record in self.records.values():
            kind = record.public["exposureType"]
            types[kind] = types.get(kind, 0) + 1
        return {
            "rows_seen": self.rows_seen,
            "valid_rows": self.valid_rows,
            "unique_records": len(self.records),
            "duplicate_rows": self.duplicate_rows,
            "invalid_rows": len(self.invalid_rows),
            "unique_media": len(media),
            "exposure_types": dict(sorted(types.items())),
        }


def _find_header(ws) -> tuple[int, dict[int, str], list[int]] | None:
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        mapped: dict[int, str] = {}
        for col_index, value in enumerate(row):
            label = _text(value)
            if label in HEADER_ALIASES:
                mapped[col_index] = HEADER_ALIASES[label]
        if {"date", "media", "title", "link"}.issubset(set(mapped.values())):
            # 歷史檔案的內部備註欄可能沒有欄名（2026/08 的第 F 欄就是
            # 這種情況）。所有未映射欄位都視為私密備註候選；實際每列只
            # 收集非空值，因此空白格式欄不會產生內容。
            # read_only 模式下，部分工具產生的 xlsx 沒有 dimension 資訊，
            # ws.max_column 會是 None；退回用表頭列實際長度，避免整個匯入崩潰。
            width = max(ws.max_column or 0, len(row))
            unknown = [index for index in range(width) if index not in mapped]
            return row_number, mapped, unknown
    return None


def parse_workbooks(paths: Iterable[Path]) -> ParseReport:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - operational guidance
        raise RuntimeError("缺少 openpyxl，請先執行 pip install -r tools/requirements.txt") from exc

    report = ParseReport()
    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in workbook.worksheets:
                header = _find_header(ws)
                if not header:
                    continue
                header_row, mapped, unknown_columns = header
                for row_number, row in enumerate(
                        ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                    if not any(_text(value) for value in row):
                        continue
                    report.rows_seen += 1
                    values = {name: row[index] if index < len(row) else None
                              for index, name in mapped.items()}
                    note_parts = [_text(values.get("note"))]
                    note_parts.extend(
                        _text(row[index]) for index in unknown_columns
                        if index < len(row) and _text(row[index])
                    )
                    note = "；".join(part for part in note_parts if part)

                    exposure_date = parse_exposure_date(values.get("date"))
                    media = normalize_media(values.get("media"))
                    reporter = _text(values.get("reporter"))
                    title = normalize_title(values.get("title"))
                    raw_link = _text(values.get("link"))
                    canonical_url = canonicalize_url(raw_link)

                    missing = []
                    if not exposure_date:
                        missing.append("日期")
                    if not media:
                        missing.append("媒體")
                    if not title:
                        missing.append("新聞標題")
                    if not canonical_url:
                        missing.append("有效連結")
                    if missing:
                        report.invalid_rows.append({
                            "file": path.name,
                            "sheet": ws.title,
                            "row": row_number,
                            "reason": "缺少" + "、".join(missing),
                        })
                        continue

                    report.valid_rows += 1
                    doc_id = exposure_id(canonical_url, exposure_date, media, title)
                    source = {"file": path.name, "sheet": ws.title, "row": row_number}
                    public = {
                        "exposureDate": exposure_date,
                        "mediaName": media,
                        "reporter": reporter,
                        "title": title,
                        "link": canonical_url,
                        "canonicalUrl": canonical_url,
                        "exposureType": infer_exposure_type(canonical_url),
                        "verified": True,
                    }
                    private = {
                        "publicId": doc_id,
                        "noteInternal": note,
                        "rawMediaName": _text(values.get("media")),
                        "rawLink": raw_link,
                    }
                    existing = report.records.get(doc_id)
                    if existing:
                        report.duplicate_rows += 1
                        existing.sources.append(source)
                        if not existing.public.get("reporter") and reporter:
                            existing.public["reporter"] = reporter
                        notes = [existing.private.get("noteInternal", ""), note]
                        existing.private["noteInternal"] = "；".join(dict.fromkeys(n for n in notes if n))
                    else:
                        report.records[doc_id] = ExposureRecord(
                            doc_id=doc_id,
                            public=public,
                            private=private,
                            sources=[source],
                        )
        finally:
            workbook.close()
    return report


def build_client(project_id: str, credentials_path: str | None = None):
    if project_id != DEST_PROJECT:
        raise RuntimeError(f"拒絕寫入非目標專案：{project_id}")
    import firebase_admin
    from firebase_admin import credentials, firestore

    app_name = "media-exposure-import"
    try:
        app = firebase_admin.get_app(app_name)
    except ValueError:
        if credentials_path:
            path = Path(credentials_path)
            if not path.is_file():
                raise RuntimeError("指定的 Service Account 檔案不存在")
            try:
                payload = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                raise RuntimeError("Service Account 檔案無法讀取或格式錯誤") from exc
            if payload.get("project_id") != project_id:
                raise RuntimeError("Service Account 所屬專案與 --project 不一致")
            cred = credentials.Certificate(str(path))
            app = firebase_admin.initialize_app(cred, {"projectId": project_id}, name=app_name)
        else:
            app = firebase_admin.initialize_app(options={"projectId": project_id}, name=app_name)
    return firestore.client(app=app)


def write_records(db, records: Iterable[ExposureRecord], import_batch: str) -> int:
    records = list(records)
    written = 0
    imported_at = datetime.now(timezone.utc)
    for start in range(0, len(records), MAX_RECORDS_PER_BATCH):
        chunk = records[start:start + MAX_RECORDS_PER_BATCH]
        batch = db.batch()
        for record in chunk:
            public = dict(record.public)
            public.update({"importBatch": import_batch, "importedAt": imported_at})
            private = dict(record.private)
            private.update({
                "sources": record.sources,
                "importBatch": import_batch,
                "importedAt": imported_at,
            })
            batch.set(db.collection(PUBLIC_COLLECTION).document(record.doc_id), public, merge=True)
            batch.set(db.collection(PRIVATE_COLLECTION).document(record.doc_id), private, merge=True)
        batch.commit()
        written += len(chunk)
    return written


def _json_safe(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    return value


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="人工媒體曝光 Excel 匯入工具（預設 dry-run）")
    parser.add_argument("workbooks", nargs="+", type=Path, help="一個或多個 .xlsx 檔案")
    parser.add_argument("--report", type=Path, help="將 dry-run 摘要寫成 JSON（不含新聞內容）")
    parser.add_argument("--copy", action="store_true", help="實際寫入 Firestore")
    parser.add_argument(
        "--approve-skipping-invalid",
        type=int,
        metavar="COUNT",
        help="明確核准略過的無效列數；必須與 dry-run 實際數量完全相同",
    )
    parser.add_argument("--project", default=DEST_PROJECT)
    parser.add_argument("--credentials", default=os.getenv("FIREBASE_SERVICE_ACCOUNT"))
    parser.add_argument("--i-approve-writing-to-transcend-news-tbm", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = make_parser().parse_args(argv)
    missing = [str(path) for path in args.workbooks if not path.is_file()]
    if missing:
        print("找不到檔案：" + "、".join(missing), file=sys.stderr)
        return 2

    report = parse_workbooks(args.workbooks)
    payload = {
        "mode": "copy" if args.copy else "dry-run",
        "project": args.project,
        **report.summary(),
        "invalid_details": report.invalid_rows,
    }
    print(json.dumps(_json_safe(payload), ensure_ascii=False, indent=2))
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(
            json.dumps(_json_safe(payload), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    if not args.copy:
        return 0
    if not args.i_approve_writing_to_transcend_news_tbm:
        print("拒絕寫入：缺少明確核准旗標", file=sys.stderr)
        return 2
    if args.project != DEST_PROJECT:
        print("拒絕寫入：--project 必須是 transcend-news-tbm", file=sys.stderr)
        return 2
    if report.invalid_rows and args.approve_skipping_invalid != len(report.invalid_rows):
        print(
            "拒絕寫入：仍有無效資料列。請先修正，或用 "
            f"--approve-skipping-invalid {len(report.invalid_rows)} 明確核准略過",
            file=sys.stderr,
        )
        return 1

    db = build_client(args.project, args.credentials)
    batch_name = "media-exposure-" + datetime.now(TAIPEI).strftime("%Y%m%d-%H%M%S")
    written = write_records(db, report.records.values(), batch_name)
    print(f"完成：寫入 {written} 筆公開資料與 {written} 筆私密來源資料")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
