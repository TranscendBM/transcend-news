import datetime
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / 'tools'))

import import_media_exposure as exposure  # noqa: E402


class _FakeDoc:
    def __init__(self, collection, doc_id):
        self.collection, self.id = collection, doc_id


class _FakeCollection:
    def __init__(self, name):
        self.name = name

    def document(self, doc_id):
        return _FakeDoc(self.name, doc_id)


class _FakeBatch:
    def __init__(self, db):
        self.db = db
        self.pending = []

    def set(self, ref, data, merge=False):
        self.pending.append((ref.collection, ref.id, dict(data), merge))

    def commit(self):
        self.db.commits.append(self.pending)


class _FakeDB:
    def __init__(self):
        self.commits = []

    def collection(self, name):
        return _FakeCollection(name)

    def batch(self):
        return _FakeBatch(self)


def _write_workbook(path, rows, include_note=True):
    wb = Workbook()
    ws = wb.active
    ws.title = '08'
    headers = ['日期', '媒體', '記者', '新聞標題', '連結']
    if include_note:
        headers.append('備註')
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


class TestNormalization(unittest.TestCase):
    def test_canonical_url_strips_tracking_and_normalizes_missing_scheme(self):
        self.assertEqual(
            exposure.canonicalize_url(' news.google.com/x?a=1&utm_source=mail#top '),
            'https://news.google.com/x?a=1',
        )

    def test_canonical_url_rejects_credentials_and_internal_whitespace(self):
        self.assertEqual(exposure.canonicalize_url('https://u:p@example.com/x'), '')
        self.assertEqual(exposure.canonicalize_url('https://exa mple.com/x'), '')

    def test_media_aliases_are_stable(self):
        self.assertEqual(exposure.normalize_media('時報新聞'), '時報資訊')
        self.assertEqual(exposure.normalize_media('工商'), '工商時報')
        self.assertEqual(exposure.normalize_media('鉅亨網新聞中心'), '鉅亨網')
        self.assertEqual(exposure.normalize_media('財訊新聞'), '財訊快報')
        self.assertEqual(exposure.normalize_media('DIGITIMES'), '電子時報')
        self.assertEqual(exposure.normalize_media('精實財經'), 'MoneyDJ理財網')
        self.assertEqual(exposure.normalize_media('中國時報'), '時報資訊')


class TestWorkbookParsing(unittest.TestCase):
    def test_parses_dedupes_and_keeps_internal_note_out_of_public_record(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Media exposure_08.xlsx'
            _write_workbook(path, [
                [datetime.date(2026, 8, 3), '工商', '王記者', '創見新品',
                 'https://example.com/story?utm_source=x', '電訪Paul'],
                [datetime.date(2026, 8, 3), '工商時報', '', '創見新品（重複）',
                 'https://example.com/story', '採訪Peter'],
            ])

            report = exposure.parse_workbooks([path])

        self.assertEqual(report.rows_seen, 2)
        self.assertEqual(report.valid_rows, 2)
        self.assertEqual(report.duplicate_rows, 1)
        self.assertEqual(len(report.records), 1)
        record = next(iter(report.records.values()))
        self.assertEqual(record.public['mediaName'], '工商時報')
        self.assertNotIn('noteInternal', record.public)
        self.assertNotIn('sourceFile', record.public)
        self.assertEqual(record.private['noteInternal'], '電訪Paul；採訪Peter')
        self.assertEqual(len(record.sources), 2)

    def test_reports_rows_missing_required_business_fields_without_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'bad.xlsx'
            _write_workbook(path, [[None, None, None, None, 'https://youtu.be/abc', None]])
            report = exposure.parse_workbooks([path])

        self.assertEqual(len(report.records), 0)
        self.assertEqual(report.invalid_rows, [{
            'file': 'bad.xlsx', 'sheet': '08', 'row': 2,
            'reason': '缺少日期、媒體、新聞標題',
        }])

    def test_blank_sheets_are_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'multi.xlsx'
            _write_workbook(path, [[datetime.date(2026, 8, 3), '經濟日報', '',
                                    '創見新品', 'https://example.com/a', '']])
            from openpyxl import load_workbook
            wb = load_workbook(path)
            wb.create_sheet('工作表2')
            wb.save(path)
            wb.close()
            report = exposure.parse_workbooks([path])
        self.assertEqual(report.summary()['unique_records'], 1)

    def test_unlabeled_extra_column_is_preserved_as_private_note(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'unlabeled-note.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = '08'
            ws.append(['日期', '媒體', '記者', '新聞標題', '連結', None])
            ws.append([datetime.date(2026, 8, 5), '非凡新聞', '洪芷茵',
                       '華強北 DRAM 新聞', 'https://example.com/a', '電訪Paul'])
            wb.save(path)
            wb.close()
            report = exposure.parse_workbooks([path])

        record = next(iter(report.records.values()))
        self.assertEqual(record.private['noteInternal'], '電訪Paul')
        self.assertNotIn('電訪Paul', json.dumps(record.public, default=str, ensure_ascii=False))


class TestWriteSafety(unittest.TestCase):
    def test_write_splits_public_and_private_payloads(self):
        when = datetime.datetime(2026, 8, 3, tzinfo=exposure.TAIPEI)
        record = exposure.ExposureRecord(
            doc_id='exp_1',
            public={'exposureDate': when, 'mediaName': '經濟日報', 'reporter': '',
                    'title': '創見新品', 'link': 'https://example.com/a',
                    'canonicalUrl': 'https://example.com/a', 'exposureType': 'online',
                    'verified': True},
            private={'publicId': 'exp_1', 'noteInternal': '電訪Paul'},
            sources=[{'file': 'source.xlsx', 'sheet': '08', 'row': 2}],
        )
        db = _FakeDB()
        self.assertEqual(exposure.write_records(db, [record], 'batch-1'), 1)
        writes = db.commits[0]
        self.assertEqual([item[0] for item in writes],
                         [exposure.PUBLIC_COLLECTION, exposure.PRIVATE_COLLECTION])
        public_payload = writes[0][2]
        private_payload = writes[1][2]
        self.assertNotIn('noteInternal', public_payload)
        self.assertNotIn('sources', public_payload)
        self.assertEqual(private_payload['noteInternal'], '電訪Paul')

    def test_cli_defaults_to_dry_run_and_never_builds_firestore_client(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'source.xlsx'
            _write_workbook(path, [[datetime.date(2026, 8, 3), '經濟日報', '',
                                    '創見新品', 'https://example.com/a', '']])
            with mock.patch.object(exposure, 'build_client') as build_client:
                code = exposure.main([str(path)])
        self.assertEqual(code, 0)
        build_client.assert_not_called()

    def test_copy_requires_explicit_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'source.xlsx'
            _write_workbook(path, [[datetime.date(2026, 8, 3), '經濟日報', '',
                                    '創見新品', 'https://example.com/a', '']])
            with mock.patch.object(exposure, 'build_client') as build_client:
                code = exposure.main([str(path), '--copy'])
        self.assertEqual(code, 2)
        build_client.assert_not_called()

    def test_copy_with_invalid_rows_requires_exact_skip_count(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'source.xlsx'
            _write_workbook(path, [
                [datetime.date(2026, 8, 3), '經濟日報', '', '創見新品',
                 'https://example.com/a', ''],
                [None, None, None, None, 'https://youtu.be/abc', None],
            ])
            approved = '--i-approve-writing-to-transcend-news-tbm'
            with mock.patch.object(exposure, 'build_client') as build_client:
                wrong = exposure.main([
                    str(path), '--copy', approved, '--approve-skipping-invalid', '2',
                ])
                self.assertEqual(wrong, 1)
                build_client.assert_not_called()

    def test_exact_invalid_skip_count_allows_only_valid_records_to_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'source.xlsx'
            _write_workbook(path, [
                [datetime.date(2026, 8, 3), '經濟日報', '', '創見新品',
                 'https://example.com/a', ''],
                [None, None, None, None, 'https://youtu.be/abc', None],
            ])
            db = _FakeDB()
            with mock.patch.object(exposure, 'build_client', return_value=db):
                code = exposure.main([
                    str(path), '--copy', '--i-approve-writing-to-transcend-news-tbm',
                    '--approve-skipping-invalid', '1',
                ])
        self.assertEqual(code, 0)
        self.assertEqual(len(db.commits), 1)
        self.assertEqual(len(db.commits[0]), 2)  # one public + one private

    def test_json_report_contains_counts_but_no_titles_or_notes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'source.xlsx'
            report_path = Path(tmp) / 'report.json'
            _write_workbook(path, [[datetime.date(2026, 8, 3), '經濟日報', '',
                                    '不應出現在報告中的標題', 'https://example.com/a',
                                    '不應出現在報告中的備註']])
            self.assertEqual(exposure.main([str(path), '--report', str(report_path)]), 0)
            raw = report_path.read_text(encoding='utf-8')
        payload = json.loads(raw)
        self.assertEqual(payload['unique_records'], 1)
        self.assertNotIn('不應出現在報告中的標題', raw)
        self.assertNotIn('不應出現在報告中的備註', raw)


if __name__ == '__main__':
    unittest.main()
